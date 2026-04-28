"""文档解析器：按 MIME 类型选择解析策略。"""

from __future__ import annotations

import io
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Protocol

from loguru import logger


class OcrEngine(Protocol):
    """扫描件 OCR 引擎（可选依赖，默认不启用）。"""

    def extract_text(self, pdf_bytes: bytes) -> str:
        ...


@dataclass
class ParsedDocument:
    """解析后的纯文本与元数据。"""

    text: str
    mime_type: str
    meta: dict[str, str]


class _TextExtractor(HTMLParser):
    """提取 HTML 可见文本，跳过 script/style。"""

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in ("script", "style"):
            self._skip += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style") and self._skip:
            self._skip -= 1

    def handle_data(self, data: str) -> None:
        if self._skip:
            return
        text = data.strip()
        if text:
            self.parts.append(text)


class DocumentParser:
    """文档解析器：支持 txt / md / html / docx / pdf。"""

    def __init__(
        self,
        max_chars: int = 500_000,
        ocr_engine: OcrEngine | None = None,
    ) -> None:
        self._max_chars = max_chars
        self._ocr_engine = ocr_engine

    def parse_file(self, path: Path, mime_type: str | None = None) -> ParsedDocument:
        """从文件路径解析。"""
        return self.parse_bytes(path.read_bytes(), path.name, mime_type)

    def parse_bytes(self, data: bytes, filename: str, mime_type: str | None) -> ParsedDocument:
        """从内存字节解析（txt / md / html / docx / pdf）。"""
        mime = mime_type or "application/octet-stream"
        lower = filename.lower()

        if lower.endswith((".md", ".markdown")):
            text = data.decode("utf-8", errors="ignore")
            return ParsedDocument(
                text=text[: self._max_chars],
                mime_type=mime,
                meta={"format": "markdown"},
            )

        if lower.endswith((".html", ".htm")) or mime == "text/html":
            return self._parse_html(data, mime)

        if lower.endswith(".docx"):
            return self._parse_docx(data, mime)

        if lower.endswith(".txt") or (mime.startswith("text/") and mime != "text/html"):
            text = data.decode("utf-8", errors="ignore")
            return ParsedDocument(text=text[: self._max_chars], mime_type=mime, meta={})

        if lower.endswith(".pdf") or mime == "application/pdf":
            return self._parse_pdf(data, mime)

        return ParsedDocument(text="", mime_type=mime, meta={"warning": "unsupported"})

    def _parse_html(self, data: bytes, mime: str) -> ParsedDocument:
        extractor = _TextExtractor()
        extractor.feed(data.decode("utf-8", errors="ignore"))
        text = "\n".join(extractor.parts)[: self._max_chars]
        return ParsedDocument(text=text, mime_type=mime, meta={"format": "html"})

    def _parse_docx(self, data: bytes, mime: str) -> ParsedDocument:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("未安装 python-docx，无法解析 .docx") from exc

        doc = DocxDocument(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)[: self._max_chars]
        return ParsedDocument(text=text, mime_type=mime, meta={"format": "docx"})

    def _parse_pdf(self, data: bytes, mime: str) -> ParsedDocument:
        """使用 pypdf 抽取文本。"""
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            logger.exception("未安装 pypdf: {}", exc)
            raise

        reader = PdfReader(io.BytesIO(data))
        parts: list[str] = []
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception as exc:
                logger.warning("单页 PDF 抽取失败: {}", exc)

        text = "\n".join(parts)[: self._max_chars]
        meta = {"pages": str(len(reader.pages))}
        if not text.strip():
            if self._ocr_engine is not None:
                try:
                    ocr_text = self._ocr_engine.extract_text(data)
                    if ocr_text.strip():
                        text = ocr_text[: self._max_chars]
                        meta["ocr"] = type(self._ocr_engine).__name__
                    else:
                        meta["ocr"] = "empty"
                except Exception as exc:  # noqa: BLE001
                    logger.warning("OCR 抽取失败: {}", exc)
                    meta["ocr"] = "error"
            if not text.strip():
                meta.setdefault("ocr", "unavailable")
        return ParsedDocument(text=text, mime_type=mime, meta=meta)

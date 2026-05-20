# -*- coding: utf-8 -*-
"""多格式解析：Markdown / HTML / DOCX。"""

from __future__ import annotations

import io

from app.etl.parser import DocumentParser


def test_markdown_is_parsed_as_text() -> None:
    document = DocumentParser().parse_bytes(
        b"# Title\n\nBody text", "note.md", "text/markdown"
    )

    assert "Title" in document.text
    assert "Body text" in document.text
    assert document.meta.get("format") == "markdown"


def test_html_strips_tags_and_scripts() -> None:
    html = (
        b"<html><body><h1>Heading</h1>"
        b"<script>var secret = 1;</script>"
        b"<p>Body text</p></body></html>"
    )

    document = DocumentParser().parse_bytes(html, "page.html", "text/html")

    assert "Heading" in document.text
    assert "Body text" in document.text
    assert "secret" not in document.text


class FakeOcr:
    def extract_text(self, pdf_bytes: bytes) -> str:
        return "OCR 识别文本"


def make_blank_pdf() -> bytes:
    from pypdf import PdfWriter

    buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(buffer)
    return buffer.getvalue()


def test_scanned_pdf_uses_ocr_engine() -> None:
    document = DocumentParser(ocr_engine=FakeOcr()).parse_bytes(
        make_blank_pdf(), "scan.pdf", "application/pdf"
    )

    assert "OCR 识别文本" in document.text
    assert document.meta["ocr"] == "FakeOcr"


def test_build_ocr_engine_disabled_returns_none() -> None:
    from app.etl.ocr import build_ocr_engine

    assert build_ocr_engine(False) is None


def test_scanned_pdf_without_ocr_reports_unavailable() -> None:
    document = DocumentParser().parse_bytes(
        make_blank_pdf(), "scan.pdf", "application/pdf"
    )

    assert document.text == ""
    assert document.meta["ocr"] == "unavailable"


def test_docx_paragraphs_are_extracted() -> None:
    from docx import Document as DocxDocument

    buffer = io.BytesIO()
    docx = DocxDocument()
    docx.add_paragraph("第一段内容")
    docx.add_paragraph("第二段内容")
    docx.save(buffer)

    document = DocumentParser().parse_bytes(
        buffer.getvalue(),
        "report.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert "第一段内容" in document.text
    assert "第二段内容" in document.text
